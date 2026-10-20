"""Excel Generator: export Composer Mapper Markdown to an operation planning workbook.

Requires Python 3.10+ and openpyxl: python -m pip install openpyxl
Usage: python "Excel Generator.py" "path/to/mapper-output"
       python "Excel Generator.py" --vault "path/to/mapper-output" --output report.xlsx --include-relationships

This portable snapshot preserves Snapir's operation discovery and wave calculation.
Copy this file alone; no Snapir installation or Node.js runtime is required.
Supports legacy vaults and scoped channel/operation vaults with simplified resource paths.
Pass the vault root, not its channels or operations folder. Scoped vaults must include
vault-manifest.json and all listed Markdown notes; original Composer sources are not needed.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass, field
import math
import sys


def write_workbook(graph, waves, output, relationships=False):
    """Write the existing migration report views as a filterable Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, title, context, headers, rows, widths, table_name):
        item = workbook.create_sheet(name)
        item.sheet_view.showGridLines = False
        item.cell(2, 1, title)
        item.cell(3, 1, context)
        for column, value in enumerate(headers, 1):
            item.cell(5, column, value)
        for row_number, values in enumerate(rows, 6):
            for column, value in enumerate(values, 1):
                cell = item.cell(row_number, column, value)
                if isinstance(value, str):
                    cell.data_type = 's'
        end = max(6, 5 + len(rows))
        for row in item.iter_rows(min_row=1, max_row=end, max_col=len(headers)):
            for cell in row:
                cell.font = Font(name='Arial', size=11, color='243247')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
        item['A2'].font = Font(name='Arial', size=16, bold=True, color='243247')
        item['A3'].font = Font(name='Arial', size=10, italic=True, color='58677B')
        item['A3'].alignment = Alignment(wrap_text=False)
        for cell in item[5]:
            cell.fill = PatternFill('solid', fgColor='234568')
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        item.row_dimensions[5].height = 32
        for column, width in enumerate(widths, 1):
            item.column_dimensions[get_column_letter(column)].width = width
        for row_number, values in enumerate(rows, 6):
            lines = max(sum(max(1, math.ceil(len(line) / max(1, widths[i] - 2)))
                            for line in str(value or '').split('\n'))
                        for i, value in enumerate(values))
            item.row_dimensions[row_number].height = min(409, max(18, lines * 15))
        table = Table(displayName=table_name, ref=f'A5:{get_column_letter(len(headers))}{end}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        item.add_table(table)
        item.freeze_panes = 'A6'
        return item

    ordered = sorted(graph['operations'], key=lambda key: (waves[key]['wave'] if waves[key]['wave'] is not None else math.inf, key))
    plan = sheet('Migration Plan', 'Operation implementation order',
        'Dependencies first. Cycles require coordinated implementation; blank waves have unresolved dependencies.',
        ['Build wave', 'Channel', 'Operation', 'Implement after (direct)', 'Cycle group', 'Issue'],
        [[waves[key]['wave'], graph['operations'][key]['channel'], graph['operations'][key]['operation_id'],
          '\n'.join(waves[key]['prerequisites']) or 'None', waves[key]['cycle_group'] or 'None', waves[key]['issue']] for key in ordered],
        [15, 18, 32, 62, 45, 45], 'MigrationPlanTable')
    for row in plan.iter_rows(min_row=6, max_row=max(6, len(ordered) + 5), max_col=1):
        row[0].number_format = '0'
    operations = sheet('Operations', 'Operation inventory',
        'Counts are unique direct operation connections, not step counts or implementation effort.',
        ['Channel', 'Operation', 'Calls count', 'Callers count', 'Build wave', 'Cycle group'],
        [[graph['operations'][key]['channel'], graph['operations'][key]['operation_id'], waves[key]['calls_count'],
          waves[key]['callers_count'], waves[key]['wave'], waves[key]['cycle_group'] or 'None'] for key in sorted(graph['operations'])],
        [18, 32, 17, 17, 16, 45], 'OperationsTable')
    for row in operations.iter_rows(min_row=6, max_row=max(6, len(ordered) + 5), min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0'
    if relationships:
        sheet('Relationships', 'Operation call evidence',
            'Filter Caller or Called operation. Multiple call sites are retained as evidence.',
            ['Caller', 'Called operation', 'Invoking step', 'Mapper note'],
            [[call['caller'], call['callee'], call['step'], call['note']] for call in graph['calls']],
            [42, 42, 32, 62], 'RelationshipsTable')
    if graph['issues']:
        sheet('Issues', 'Unresolved operation relationships',
            'Affected operations and their callers have no assigned build wave.',
            ['Operation', 'Step', 'Issue'],
            [[issue['operation'], issue['step'], issue['detail']] for issue in graph['issues']],
            [42, 45, 90], 'IssuesTable')
    workbook.save(output)


def main(argv=None):
    """Export a Mapper output directory without importing any Snapir modules."""
    parser = argparse.ArgumentParser(prog='Excel Generator', description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('mapper_output', nargs='?', type=Path, help='Mapper vault root; scoped vaults must include vault-manifest.json')
    parser.add_argument('--vault', type=Path, help='Alias for the Mapper output directory')
    parser.add_argument('--output', type=Path, help='Destination .xlsx; defaults beside the Mapper output folder')
    parser.add_argument('--include-relationships', action='store_true')
    args = parser.parse_args(argv)
    if (args.mapper_output is None) == (args.vault is None):
        parser.error('provide one Mapper output path, either positionally or with --vault')
    vault = (args.mapper_output or args.vault).resolve()
    output = args.output or vault.parent / 'operation-migration.xlsx'
    if output.suffix.lower() != '.xlsx':
        parser.error('--output must end in .xlsx')
    try:
        graph = discover_operation_calls(vault)
        waves = operation_waves(graph)
        output.parent.mkdir(parents=True, exist_ok=True)
        write_workbook(graph, waves, output, args.include_relationships)
    except ModuleNotFoundError as exc:
        if exc.name != 'openpyxl':
            raise
        parser.exit(1, 'Excel Generator requires openpyxl. Install it with: python -m pip install openpyxl\n')
    except (OSError, ValueError) as exc:
        parser.exit(1, f'Excel Generator: {exc}\n')
    print(f"Created {output.resolve()}: {len(graph['operations'])} operations, {len(graph['calls'])} call sites, {len(graph['issues'])} issues.")
    return 0

@dataclass(kw_only=True)
class ComposerOperation:
    id: str
    channel: str
    source_file: str
    source_hash: str
    operation_context: str | None = None
    host_key: str | None = None
    operation_fields: str | None = None
    write_to_log: str | None = None
    write_to_ofec_stat: str | None = None
    is_selective_journalising: str | None = None
    ini_values: list[str] = field(default_factory=list)
    steps: list[str] = field(default_factory=list)
    ref_formats: dict[str, str] = field(default_factory=dict)
    inline_format_ids: list[str] = field(default_factory=list)
    conversion_status: str

@dataclass(kw_only=True)
class ComposerOpStep:
    id: str
    parent_operation_id: str
    impl_class: str | None = None
    source_file: str
    source_hash: str
    only_for: str | None = None
    transitions: dict[str, str] = field(default_factory=dict)
    return_body_switches: dict[str, str] = field(default_factory=dict)
    parameters: dict[str, str] = field(default_factory=dict)
    conversion_status: str

# Bundled from source/tools/composer_mapper_tools/frontmatter_parser.py
import re
from typing import Any

def parse_frontmatter(md_text: str) -> dict[str, Any]:
    match = re.search('(?s)(?:^|\\n)---\\n(.*?)\\n---(?:\\n|$)', md_text)
    if match is None:
        return {}
    return _parse_yaml_subset(match.group(1))

def parse_section(md_text: str, section_heading: str) -> str:
    pattern = re.compile(f'(?ms)^##\\s+{re.escape(section_heading)}\\s*\\n(.*?)(?=^##\\s+|\\Z)')
    match = pattern.search(md_text)
    if match is None:
        return ''
    return match.group(1).strip()

def _parse_yaml_subset(text: str) -> dict[str, Any]:
    result: dict[str, Any] = {}
    current_key: str | None = None
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if current_key and line.startswith('  - '):
            result[current_key].append(_coerce_scalar(line[4:].strip()))
            continue
        current_key = None
        if ':' not in line:
            continue
        key, value = line.split(':', 1)
        key = key.strip()
        value = value.strip()
        if value == '[]':
            result[key] = []
        elif value == '':
            result[key] = []
            current_key = key
        else:
            result[key] = _coerce_scalar(value)
    return result

def _coerce_scalar(value: str) -> Any:
    value = value.strip().strip('"').strip("'")
    if value.lower() == 'true':
        return True
    if value.lower() == 'false':
        return False
    return value

def clean_markdown_value(value: str) -> str:
    value = value.strip()
    value = re.sub('\\[([^\\]]+)\\]\\([^)]+\\)', '\\1', value)
    value = value.replace('`', '')
    return value.strip()

def parse_bullet_values(section_text: str) -> list[str]:
    values = []
    for line in section_text.splitlines():
        stripped = line.strip()
        if stripped.startswith('- '):
            value = clean_markdown_value(stripped[2:])
            if value and value.lower() != 'none':
                values.append(value)
    return values

def parse_key_value_bullets(section_text: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for item in parse_bullet_values(section_text):
        if ':' in item:
            key, value = item.split(':', 1)
            values[key.strip()] = value.strip()
        elif '->' in item:
            key, value = item.split('->', 1)
            values[key.strip()] = value.strip()
    return values

# Bundled from source/tools/composer_mapper_tools/operation_parser.py
def parse_operation_note(md_text: str, source_path: str) -> ComposerOperation:
    frontmatter = parse_frontmatter(md_text)
    ref_formats = {}
    request = _first_bullet(parse_section(md_text, 'Request Format'))
    reply = _first_bullet(parse_section(md_text, 'Reply Format'))
    if request:
        ref_formats['csRequestFormat'] = request
    if reply:
        ref_formats['csReplyFormat'] = reply
    ref_formats.update(parse_key_value_bullets(parse_section(md_text, 'Other Formats')))
    operation_attrs = parse_key_value_bullets(parse_section(md_text, 'Operation Attributes'))
    context = _first_bullet(parse_section(md_text, 'Context'))
    if context and context.lower().startswith('unresolved:'):
        context = context.split(':', 1)[1].strip()
    return ComposerOperation(id=str(frontmatter.get('entity_id', '')), channel=str(frontmatter.get('channel', '')), source_file=str(frontmatter.get('source_file', '')), source_hash=str(frontmatter.get('source_hash', '')), operation_context=context, host_key=frontmatter.get('host_key'), operation_fields=operation_attrs.get('operationFields'), write_to_log=operation_attrs.get('writeToLog'), write_to_ofec_stat=operation_attrs.get('writeToOfecStat'), is_selective_journalising=operation_attrs.get('isSelectiveJournalising'), ini_values=parse_bullet_values(parse_section(md_text, 'Initial Values (iniValue)')), steps=[_strip_step_suffix(step) for step in parse_bullet_values(parse_section(md_text, 'Steps'))], ref_formats=ref_formats, inline_format_ids=[value.removeprefix('inline:').strip() for value in parse_bullet_values(parse_section(md_text, 'Other Formats')) if value.startswith('inline:')], conversion_status=str(frontmatter.get('conversion_status', 'not_started')))

def _first_bullet(section_text: str) -> str | None:
    values = parse_bullet_values(section_text)
    return values[0] if values else None

def _strip_step_suffix(value: str) -> str:
    return clean_markdown_value(value.split(' *(', 1)[0])

# Bundled from source/tools/composer_mapper_tools/opstep_parser.py
import re

def parse_opstep_note(md_text: str, source_path: str) -> ComposerOpStep:
    frontmatter = parse_frontmatter(md_text)
    implementation = _parse_impl_class(parse_section(md_text, 'Implementation'))
    return ComposerOpStep(id=str(frontmatter.get('entity_id', '')), parent_operation_id=_first(frontmatter.get('used_by_operations', [])) or '', impl_class=implementation, source_file=str(frontmatter.get('source_file', '')), source_hash=str(frontmatter.get('source_hash', '')), only_for=frontmatter.get('onlyFor'), transitions=_parse_transitions(parse_section(md_text, 'Transitions')), return_body_switches=_parse_return_body_switches(parse_section(md_text, 'Return Body Switches')), parameters=parse_key_value_bullets(parse_section(md_text, 'Parameters')), conversion_status=str(frontmatter.get('conversion_status', 'not_started')))

def _parse_impl_class(section: str) -> str | None:
    for line in section.splitlines():
        if line.strip().startswith('- '):
            value = clean_markdown_value(line.strip()[2:])
            if value.lower() == 'unresolved':
                return None
            if value.lower().startswith('unresolved:'):
                return value.split(':', 1)[1].strip()
            return value
    return None

def _parse_transitions(section: str) -> dict[str, str]:
    transitions: dict[str, str] = {}
    for line in section.splitlines():
        value = clean_markdown_value(line.strip().removeprefix('- '))
        if ' -> ' in value:
            key, target = value.split(' -> ', 1)
            transitions[key.strip()] = target.removeprefix('unresolved ').strip()
    return transitions

def _parse_return_body_switches(section: str) -> dict[str, str]:
    switches: dict[str, str] = {}
    for line in section.splitlines():
        if not line.strip().startswith('- '):
            continue
        value = clean_markdown_value(line.strip()[2:])
        match = re.search('RC\\s+(\\S+)\\s+->\\s+switch reply body to\\s+(.+)', value)
        if match:
            switches[match.group(1)] = match.group(2).strip()
    return switches

def _first(value: object) -> str | None:
    if isinstance(value, list) and value:
        return str(value[0])
    return None

# Bundled from source/tools/composer_mapper_tools/channel_registry_parser.py
from pathlib import Path

def parse_channel_operation_registry(registry_md_path: str) -> list[tuple[str, str, str]]:
    path = Path(registry_md_path)
    if not path.exists():
        return []
    rows = []
    for line in path.read_text().splitlines():
        stripped = line.strip()
        if not stripped.startswith('|') or '---' in stripped or 'Operation' in stripped:
            continue
        cells = [clean_markdown_value(cell.strip()) for cell in stripped.strip('|').split('|')]
        if len(cells) >= 3:
            rows.append((cells[0], cells[1], cells[2]))
    return rows

# Bundled from source/tools/composer_mapper_tools/vault_manifest.py
"""Validate scoped vault identity and enumerate notes without guessing folder depth."""
import hashlib
import json
from pathlib import Path

def read_vault_manifest(vault):
    """Return a validated schema-two manifest, or None for a legacy vault.

    Incomplete generation, unknown schemas, conflicting views and escaped paths are errors.
    This reader uses only Python's standard library so standalone exporters can bundle it.
    """
    root = Path(vault).resolve()
    path = root / 'vault-manifest.json'
    if (root / 'vault-generation.json').exists():
        raise ValueError('Incomplete scoped vault generation: ' + str(root))
    if not path.exists():
        if (root / 'channels').exists() and list((root / 'channels').glob('*/operations/*/operation.md')):
            raise ValueError('Scoped vault is missing vault-manifest.json')
        return None
    manifest = json.loads(path.read_text(encoding='utf-8-sig'))
    if not isinstance(manifest, dict):
        raise ValueError('Invalid scoped vault manifest')
    if manifest.get('vault_schema_version') != 2 or manifest.get('identity_version') != 1:
        raise ValueError('Unsupported vault schema or identity version')
    if manifest.get('complete') is not True or not manifest.get('workspace_id'):
        raise ValueError('Incomplete scoped vault manifest')
    entries = manifest.get('entities')
    if not isinstance(entries, list):
        raise ValueError('Missing scoped entity inventory')
    paths, usages, definitions = ({}, set(), {})
    for entry in entries:
        required = ('kind', 'entity_id', 'definition_id', 'usage_id', 'note', 'source_relative_path', 'source_selector', 'source_hash', 'note_hash')
        if not isinstance(entry, dict) or any((not isinstance(entry.get(key), str) or not entry[key] for key in required)):
            raise ValueError('Invalid scoped entity identity')
        relative = entry['note']
        candidate = (root / relative).resolve()
        if '\\' in relative or Path(relative).is_absolute() or (not candidate.is_relative_to(root)) or (not candidate.is_file()):
            raise ValueError('Missing or unsafe scoped note path: ' + relative)
        if relative.casefold() in paths or entry['usage_id'] in usages:
            raise ValueError('Duplicate scoped note path or usage identity: ' + relative)
        if hashlib.sha256(candidate.read_bytes()).hexdigest() != entry['note_hash']:
            raise ValueError('Scoped note differs from completed manifest: ' + relative)
        metadata = parse_frontmatter(candidate.read_text(encoding='utf-8-sig'))
        for field, key in (('entity_id', 'entity_id'), ('qualified_entity_id', 'definition_id'), ('usage_id', 'usage_id')):
            if str(metadata.get(field, '')) != entry[key]:
                raise ValueError('Scoped note and manifest identity disagree: ' + relative)
        if str(metadata.get('operation_identity') or '') != entry.get('operation_identity', ''):
            raise ValueError('Scoped note belongs to another operation view: ' + relative)
        coordinates = tuple((entry[key] for key in ('kind', 'entity_id', 'source_relative_path', 'source_selector', 'source_hash')))
        previous = definitions.setdefault(entry['definition_id'], coordinates)
        if previous != coordinates:
            raise ValueError('Conflicting views of source definition: ' + entry['definition_id'])
        paths[relative.casefold()] = entry
        usages.add(entry['usage_id'])
    for entry in entries:
        if entry['kind'] != 'operation' or not entry.get('operation_id'):
            continue
        if entry['operation_id'] != entry['entity_id'] or not entry.get('channel') or (not entry.get('operation_identity')):
            raise ValueError('Invalid effective operation identity: ' + entry['note'])
        metadata = parse_frontmatter((root / entry['note']).read_text(encoding='utf-8-sig'))
        if str(metadata.get('channel', '')) != entry['channel']:
            raise ValueError('Scoped operation channel and manifest disagree: ' + entry['note'])
        for reference in entry.get('references', {}).values():
            target = paths.get(reference.casefold())
            if target is None or target.get('operation_identity') != entry['operation_identity']:
                raise ValueError('Reference escapes owning operation view: ' + reference)
    return manifest

def operation_note_paths(vault, manifest=None):
    """Enumerate effective operation notes from a scoped manifest or legacy folders."""
    root = Path(vault).resolve()
    manifest = read_vault_manifest(root) if manifest is None else manifest
    if manifest is not None:
        return [root / entry['note'] for entry in manifest['entities'] if entry['kind'] == 'operation' and entry.get('operation_id')]
    if not (root / 'operations').is_dir():
        raise ValueError('Missing operations folder: ' + str(root))
    return sorted((root / 'operations').rglob('*.md'))

# Bundled from source/tools/operation_call_tools/discovery.py
"""Resolve CCComplexOpStep targets from Composer Mapper Markdown evidence."""
import hashlib
import re
from pathlib import Path
from urllib.parse import unquote
COMPLEX_CLASS = 'com.ibm.il.dse.cc.cut.CCComplexOpStep'

def _links(section):
    """Read both standard Markdown links and Obsidian links."""
    return [unquote(match[0] or match[1].split('|', 1)[0]) for match in re.findall('\\[[^\\]]*\\]\\(([^)]+)\\)|\\[\\[([^\\]]+)\\]\\]', section)]

def _note_path(root, owner, target):
    """Resolve Mapper vault-root links or ordinary relative note links safely."""
    target = target.split('#', 1)[0]
    if not target.endswith('.md'):
        target += '.md'
    for candidate in (root / target, owner.parent / target):
        path = candidate.resolve()
        if path.is_relative_to(root) and path.is_file():
            return path
    raise ValueError('Missing linked step note: ' + target)

def discover_operation_calls(vault):
    """Read a vault without mutation; return channel-qualified operations, calls and issues."""
    root = Path(vault).resolve()
    manifest = read_vault_manifest(root)
    note_paths = operation_note_paths(root, manifest)
    registry = parse_channel_operation_registry(str(root / 'analysis/channel-operation-registry.md'))
    operations, calls, issues = ({}, [], [])
    if manifest is not None:
        issues.extend((dict(operation=item.get('operation', ''), step=item.get('step', ''), detail=item['detail']) for item in manifest.get('diagnostics', []) if item.get('kind') in {'identity', 'source_parse'}))
    documents = []
    for path in note_paths:
        text = path.read_text(encoding='utf-8-sig')
        op = parse_operation_note(text, str(path))
        channels = {op.channel} if op.channel else {channel for name, channel, _ in registry if name == op.id}
        if not op.id or not channels:
            issues.append(dict(operation=str(path.relative_to(root)), step='', detail='Operation ID or channel is missing'))
            continue
        for channel in sorted(channels):
            key = channel + '/' + op.id
            if key in operations:
                issues.append(dict(operation=key, step='', detail='Duplicate operation identity'))
                continue
            operations[key] = dict(operation_id=op.id, channel=channel, note=str(path.relative_to(root)), source_hash=op.source_hash, ref_formats=op.ref_formats)
            documents.append((key, op, path, text))
    for caller, op, path, text in documents:
        targets = _links(parse_section(text, 'Steps'))
        if op.steps and (not targets):
            issues.append(dict(operation=caller, step='', detail='Steps have no resolvable Markdown links'))
        for target in targets:
            try:
                step_path = _note_path(root, path, target)
                step_text = step_path.read_text(encoding='utf-8-sig')
                step = parse_opstep_note(step_text, str(step_path))
                raw = parse_key_value_bullets(parse_section(step_text, 'Raw Attributes'))
                impl = raw.get('implClass') or step.impl_class or ''
                if impl != COMPLEX_CLASS and (not impl.endswith('|' + COMPLEX_CLASS + ']]')):
                    continue
                channel = caller.split('/', 1)[0]
                if step.only_for and channel.upper() not in re.split('[,;\\s]+', step.only_for.strip().upper()):
                    continue
                if step.parent_operation_id and step.parent_operation_id != op.id:
                    raise ValueError('Linked step belongs to another operation; possible Mapper step-ID collision')
                name = step.parameters.get('opExecuteName') or raw.get('opExecuteName')
                if not name:
                    raise ValueError('CCComplexOpStep has no opExecuteName')
                callee = channel + '/' + name
                if callee not in operations:
                    issues.append(dict(operation=caller, step=step.id, callee=callee, note=str(step_path.relative_to(root)), parameters=step.parameters, transitions=step.transitions, evidence_hash=hashlib.sha256(step_text.encode()).hexdigest(), detail='Target is missing in caller channel: ' + callee + '; cross-channel targets require explicit resolution'))
                    continue
                calls.append(dict(caller=caller, callee=callee, step=step.id, note=str(step_path.relative_to(root)), parameters=step.parameters, transitions=step.transitions, evidence_hash=hashlib.sha256(step_text.encode()).hexdigest()))
            except (OSError, ValueError) as exc:
                issues.append(dict(operation=caller, step=target, detail=str(exc)))
    return dict(operations=operations, calls=calls, issues=issues)

# Bundled from source/tools/operation_call_tools/ordering.py
"""Calculate dependency-first waves with explicit strongly connected groups."""
import heapq
import hashlib

def operation_waves(graph):
    """Return deterministic SCC groups and waves without recursion limits on large vaults."""
    edges = {key: set() for key in graph['operations']}
    reverse = {key: set() for key in edges}
    for call in graph['calls']:
        edges[call['caller']].add(call['callee'])
        reverse[call['callee']].add(call['caller'])
    seen, finished = (set(), [])
    for start in sorted(edges):
        if start in seen:
            continue
        stack = [(start, False)]
        while stack:
            node, exiting = stack.pop()
            if exiting:
                finished.append(node)
            elif node not in seen:
                seen.add(node)
                stack.append((node, True))
                stack.extend(((other, False) for other in sorted(edges[node], reverse=True) if other not in seen))
    groups, assigned = ([], {})
    for start in reversed(finished):
        if start in assigned:
            continue
        group, pending = ([], [start])
        index = len(groups)
        assigned[start] = index
        while pending:
            node = pending.pop()
            group.append(node)
            for other in reverse[node]:
                if other not in assigned:
                    assigned[other] = index
                    pending.append(other)
        groups.append(sorted(group))
    dependencies = {i: {assigned[target] for node in nodes for target in edges[node] if assigned[target] != i} for i, nodes in enumerate(groups)}
    dependents = {i: set() for i in dependencies}
    for i, targets in dependencies.items():
        for target in targets:
            dependents[target].add(i)
    remaining = {i: len(targets) for i, targets in dependencies.items()}
    ready = [(groups[i][0], i) for i, count in remaining.items() if not count]
    heapq.heapify(ready)
    waves = {}
    while ready:
        _, i = heapq.heappop(ready)
        waves[i] = 1 + max((waves[target] for target in dependencies[i]), default=0)
        for parent in dependents[i]:
            remaining[parent] -= 1
            if remaining[parent] == 0:
                heapq.heappush(ready, (groups[parent][0], parent))
    blocked = {issue['operation'] for issue in graph['issues'] if issue['operation'] in edges}
    cyclic = {node for nodes in groups for node in nodes if len(nodes) > 1 or node in edges[node]}
    cycle_ids = {i: 'cycle-' + hashlib.sha256('\n'.join(nodes).encode()).hexdigest()[:12] for i, nodes in enumerate(groups) if any((node in cyclic for node in nodes))}
    pending = list(blocked)
    while pending:
        for parent in reverse[pending.pop()]:
            if parent not in blocked:
                blocked.add(parent)
                pending.append(parent)
    return {node: dict(wave=None if node in blocked else waves[assigned[node]], cycle_group=cycle_ids.get(assigned[node], ''), prerequisites=sorted(edges[node]), calls_count=len(edges[node]), callers_count=len(reverse[node]), issue='Unresolved dependency' if node in blocked else 'Cycle requires coordinated contracts' if node in cyclic else '') for node in sorted(edges)}

if __name__ == '__main__':
    raise SystemExit(main())
